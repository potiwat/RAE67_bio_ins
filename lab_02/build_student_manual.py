"""Build the student-facing Lab 02 manual as an editable Word document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
DESTINATION = HERE / "Lab02_Student_Manual_TH.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "5")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)


def set_cell_padding(cell, top=110, start=110, bottom=110, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        child = OxmlElement(f"w:{side}")
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")
        margins.append(child)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
    set_repeat_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_padding(cell)
            if widths:
                cell.width = Inches(widths[column_index])
            if row_index == 0:
                set_cell_shading(cell, "DCECF0")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F7FAFB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Tahoma"
                    run.font.size = Pt(9)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph()
    return table


def add_body(doc, text=""):
    paragraph = doc.add_paragraph(text)
    paragraph.style = doc.styles["Normal"]
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_command(doc, command):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(command)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    return paragraph


def add_answer_lines(doc, count=2):
    for _ in range(count):
        paragraph = doc.add_paragraph("________________________________________________________________________________")
        paragraph.paragraph_format.space_after = Pt(5)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Tahoma"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in (("Title", 19, 0, 10), ("Heading 1", 13, 15, 7), ("Heading 2", 11, 11, 5)):
    style = styles[name]
    style.font.name = "Tahoma"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Title"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

doc.add_paragraph("คู่มือนักศึกษา Lab 02 Robust Reactive Lamp", style="Title")
add_body(doc, "รายวิชา 1353410 Biological Sensing and Sensorimotor Control")
add_body(doc, "ชื่อกลุ่ม ____________________  สมาชิก _________________________________________  วันที่ ____________")
add_body(doc, "แล็บนี้ให้ทดลองใน simulator ก่อน เพื่อเปรียบเทียบ controller 4 แบบภายใต้ระยะจริงจำลองชุดเดียวกัน นักศึกษาจะเก็บข้อมูลซ้ำ วิเคราะห์ความเร็ว ความผิดพลาด และความเสถียร แล้วสรุปจากตัวเลขและกราฟ ไม่ใช้การสาธิตครั้งเดียวเป็นหลักฐาน")

doc.add_heading("เป้าหมายการเรียนรู้", level=1)
for item in (
    "อธิบายความต่างระหว่างค่าระยะจริง ค่าที่ sensor อ่าน และค่าที่ controller ใช้",
    "สร้างและเปรียบเทียบ C0 open loop, C1 single threshold, C2 filter plus hysteresis และ C3 proportional response",
    "วัด response latency, false trigger, switching count, recovery time และ command smoothness",
):
    add_bullet(doc, item)

doc.add_heading("ไฟล์ที่ใช้และข้อควรระวัง", level=1)
add_body(doc, "ทำงานในโฟลเดอร์ lab_02 ของ GitHub repo RAE67_bio_ins โดยใช้ lab02.py, student_controller.py, grade_submission.py, plot_trial.py, requirements.txt และ tests/")
add_body(doc, "แหล่งไฟล์: https://github.com/potiwat/RAE67_bio_ins/tree/main/lab_02 (repo เป็น private ต้องได้รับสิทธิ์เข้าถึงก่อน)")
add_body(doc, "ค่าระยะ threshold, gain และ command limit ในเอกสารนี้เป็น proposed course model สำหรับ simulator เท่านั้น ห้ามใช้กับหุ่นยนต์จริงก่อนตรวจ datasheet, mechanical limit, หน่วยคำสั่ง และ emergency stop")

doc.add_heading("ขั้นที่ 1 เขียนสมมติฐานก่อนรัน", level=1)
add_body(doc, "ตอบก่อนเปิดไฟล์ผล หากแก้คำตอบหลังทดลอง ให้ลงวันที่และอธิบายเหตุผล")
for text in (
    "เมื่อ noise เพิ่ม C1 จะมี switching count เปลี่ยนอย่างไร เพราะอะไร",
    "C2 จะลด false trigger โดยแลกกับ latency หรือ recovery time อย่างไร",
    "เมื่อเพิ่ม sensor delay 300 ms การตอบสนองจะเลื่อนจากเหตุการณ์จริงอย่างไร",
):
    add_body(doc, text)
    add_answer_lines(doc, 1)

doc.add_heading("ขั้นที่ 2 เตรียมโปรแกรม", level=1)
add_body(doc, "เปิด PowerShell ในโฟลเดอร์แล็บ ใช้ Python 3.10 ขึ้นไป แล้วรันคำสั่งต่อไปนี้")
for command in (
    "python -m venv .venv",
    ".\\.venv\\Scripts\\Activate.ps1",
    "python -m pip install -r requirements.txt",
    "python -m unittest discover -s tests -v",
):
    add_command(doc, command)
add_body(doc, "หาก PowerShell ไม่อนุญาตให้ activate ให้ใช้ .venv\\Scripts\\python.exe แทน python ในคำสั่งถัดไป โดยไม่ต้องเปลี่ยน execution policy ของเครื่อง")
add_body(doc, "ผลที่คาด: tests ผ่าน 12 รายการ หากไม่ผ่าน ให้จด error และแก้ก่อนเริ่มเก็บข้อมูล ชุดนี้ทดสอบตัวเชื่อม แต่ไม่ถือว่า TODO ของนักศึกษาเสร็จแล้ว")

doc.add_heading("ขั้นที่ 3 อ่านและเติม controller", level=1)
add_body(doc, "เปิด student_controller.py และทำ TODO 1 ถึง TODO 4: low pass filter, single threshold, hysteresis และ proportional command พร้อม clamp จากนั้นเทียบแนวคิดกับ lab02.py ซึ่งเป็น reference implementation")
add_body(doc, "lab02.py เชื่อม StudentController เข้ากับ experiment loop แล้ว เมธอด TODO แต่ละตัวต้องคืนค่าที่คำนวณได้ โหมด reference ใช้ตัวอย่างใน lab02.py ส่วนโหมด student ใช้ไฟล์ที่ส่งด้วย --student-file")
add_table(doc, ["Condition", "หลักการ", "ข้อมูลที่ใช้"], [
    ("C0", "ทำ sequence ตามเวลา", "ไม่ใช้ sensor"),
    ("C1", "threshold 0.35 m", "raw distance"),
    ("C2", "filter β = 0.25 และ hysteresis 0.32/0.38 m", "filtered distance"),
    ("C3", "Kp = 2.0 และ clamp [-0.6, 0.6]", "filtered distance"),
], [0.8, 3.8, 2.2])
add_body(doc, "ตรวจทิศทางสัญญาณ: ใน simulator command บวกหมายถึงตอบสนองเมื่อมือเข้าใกล้ ส่วน invalid sample ให้ safe command = 0 และบันทึกเหตุการณ์ INVALID")

doc.add_heading("ขั้นที่ 4 ทดลอง A เปรียบเทียบ C0 ถึง C3", level=1)
add_body(doc, "รัน 5 trials ต่อ condition ที่ noise σ = 0.01 m และ delay = 0 ms รวม 20 trials ทุก condition ใช้ input trajectory เดียวกัน")
add_command(doc, "python lab02.py --protocol controller-comparison --out results\\comparison")
add_body(doc, "หลังเติม TODO ครบ ให้ทดลอง controller ของกลุ่มด้วย input และ metric เดียวกัน แต่เก็บผลคนละโฟลเดอร์")
add_command(doc, "python lab02.py --controller-source student --student-file student_controller.py --protocol controller-comparison --out results\\student_comparison")
add_command(doc, "python grade_submission.py --student-file student_controller.py --out results\\student_grade.json")
add_body(doc, "grade_submission.py รายงาน PASS/FAIL ของ filter, threshold, hysteresis, proportional/clamp, invalid sample และเปรียบเทียบ trace C0–C3 กับ reference หาก TODO ยังไม่เสร็จจะคืน exit code 1 โดยยังไม่สร้างผลทดลองของโหมด student")
add_body(doc, "ตรวจว่ามี run_config.json, summary_trials.csv, summary_aggregate.csv และ raw CSV ในโฟลเดอร์ trials เปิด run_config.json เพื่อตรวจ seed, parameter, controller_source และ SHA-256 ของไฟล์นักศึกษา")
add_body(doc, "Input trajectory 20 s: 0–4 s อยู่ไกล 0.60 m; 4–8 s เข้าใกล้ถึง 0.25 m; 8–12 s ค้างใกล้พร้อม disturbance เล็กน้อย; 12–16 s เคลื่อนออก; 16–20 s อยู่ไกล")

doc.add_heading("ขั้นที่ 5 ทดลอง B ผลของ noise และ delay", level=1)
add_body(doc, "เลือก controller หนึ่งแบบจากขั้น A และระบุเหตุผล หากยังไม่เลือก ให้ใช้ C2 ตามตัวอย่าง ชุดย่อมี 5 combinations × 3 trials = 15 trials")
add_command(doc, "python lab02.py --protocol robustness-short --condition C2 --out results\\c2_short")
add_body(doc, "ชุดเต็มสำหรับงานหลังเรียนมี noise 3 ระดับ × delay 3 ระดับ × 5 trials = 45 trials")
add_command(doc, "python lab02.py --protocol robustness-full --condition C2 --out results\\c2_full")
add_body(doc, "อย่าใช้ --overwrite เว้นแต่ต้องการลบผลในโฟลเดอร์เดิมอย่างตั้งใจ ใช้ชื่อ output ใหม่เพื่อเก็บหลักฐานเดิม")

doc.add_heading("ขั้นที่ 6 สร้างกราฟและอ่านผล", level=1)
add_body(doc, "เลือก trial ตัวแทนจาก C1 และ C2 สร้างกราฟจาก raw CSV โดยคงแกนเวลาเดียวกัน")
add_command(doc, "python plot_trial.py results\\comparison\\trials\\C2_n0.01_d000_r01.csv --out results\\comparison\\C2_trial1.png")
add_body(doc, "บนกราฟให้ทำเครื่องหมายเวลาที่ true distance ข้าม 0.35 m และเวลาที่ state ตอบสนองถูกต้องครั้งแรก แยก enter กับ exit")
add_table(doc, ["Metric", "อ่านค่าอย่างไร"], [
    ("Latency enter/exit", "เวลาตอบสนองครั้งแรก ลบเวลา physical crossing"),
    ("False trigger rate", "จำนวนเข้า NEAR ผิด ÷ valid observation time รายงานครั้งต่อนาที"),
    ("Recovery time", "เวลาจาก physical crossing จนเริ่มช่วง state ถูกต้องต่อเนื่อง 0.5 s รายงานค่ามากกว่าของ enter/exit"),
    ("Extra switches", "switching count ที่เกินการเปลี่ยนตามแผน 2 ครั้ง"),
    ("Mean |Δu|", "ค่าเฉลี่ยการเปลี่ยน command ระหว่าง sample ต่อเนื่อง"),
], [2.0, 4.8])
add_body(doc, "C0 อาจเปลี่ยน state ก่อน physical crossing เพราะทำงานตามเวลา และ command ที่ไม่เคลื่อนอาจมี Mean |Δu| ต่ำแต่ไม่ใช่ controller ที่ดี จึงต้องอ่าน metric ร่วมกัน")

doc.add_heading("ใบงานบันทึกผล", level=1)
add_body(doc, "กรอกค่าจาก summary_aggregate.csv เท่านั้น หาก metric ไม่มีค่า ให้ใส่ N/A และอธิบายสาเหตุ ไม่ประมาณค่าขึ้นเอง")
add_table(doc, ["Condition", "Latency enter mean ± SD s", "False trigger ครั้ง/นาที", "Extra switches", "Mean |Δu|"], [
    ("C0", "", "", "", ""),
    ("C1", "", "", "", ""),
    ("C2", "", "", "", ""),
    ("C3", "", "", "", ""),
], [0.8, 1.6, 1.6, 1.3, 1.5])
add_body(doc, "Controller ที่เลือกทดสอบ robustness __________  เหตุผล ________________________________")
add_body(doc, "จำนวน invalid samples __________  จำนวน trial ที่ exclude __________  เหตุผล ________________")

doc.add_heading("คำถามวิเคราะห์", level=1)
for question in (
    "C1 กับ C2 ต่างกันด้าน extra switches และ latency เท่าใด อ้างตัวเลขและหน่วย",
    "noise 0.01 กับ 0.03 m ส่งผลต่อ metric ใดชัดที่สุด",
    "delay 300 ms ทำให้ latency เพิ่มใกล้ 300 ms หรือไม่ ใช้กราฟประกอบคำอธิบาย",
    "C3 ชน command limit ช่วงใด และเหตุใด smoothness อย่างเดียวไม่พอ",
):
    add_body(doc, question)
    add_answer_lines(doc, 2)

doc.add_heading("แนวทางเขียนรายงาน", level=1)
add_body(doc, "รายงาน 1–2 หน้าเริ่มด้วยคำถามวิจัยและ prediction จากนั้นระบุ protocol, seed และจำนวน trial แสดงตารางผลพร้อมหน่วย เลือกกราฟที่ชี้เหตุการณ์สำคัญ แล้วอธิบายว่าข้อมูลสนับสนุนหรือขัดกับ prediction อย่างไร ปิดด้วยข้อจำกัด เช่น simulator ไม่แทน sensor และ actuator จริง")
add_body(doc, "หากตัด trial หรือพบ invalid sample ต้องรายงานจำนวนและเหตุผลอย่างเปิดเผย ห้ามแก้ raw CSV หรือคัดเฉพาะผลที่สนับสนุนสมมติฐาน")
add_body(doc, "ผล PASS/FAIL จาก grade_submission.py ใช้ตรวจการทำงานของโค้ดเท่านั้น ไม่แทนคะแนน 20 คะแนน ผู้สอนยังต้องตรวจ block diagram, protocol, การวิเคราะห์และการตีความ")

doc.add_heading("สิ่งที่ต้องส่ง", level=1)
for item in (
    "สมมติฐานก่อนทดลองและบันทึกการแก้สมมติฐาน หากมี",
    "Block diagram ของ sensorimotor loop ที่อธิบายทางข้อมูลและ feedback",
    "student_controller.py ที่เติม TODO พร้อม source code อื่นที่แก้",
    "ผลตรวจจาก grade_submission.py และ run_config.json ของโหมด student",
    "run_config.json, raw CSV ทุก trial และ summary_aggregate.csv",
    "กราฟตัวแทนอย่างน้อย C1 และ C2 พร้อมตำแหน่ง physical crossing และ first correct response",
    "รายงาน 1–2 หน้า ระบุ condition, metric, ตัวเลขเปรียบเทียบ และข้อจำกัดของ simulator",
):
    add_bullet(doc, item)

doc.save(DESTINATION)
print(DESTINATION)
